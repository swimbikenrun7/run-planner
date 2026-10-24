"""Import a training plan into plan_builder.xlsx in three steps.

    1. python import_plan.py prompt plans/plan_builder.xlsx
         writes imports/prompt.md: paste it into Claude with the plan, save the CSV
         it returns as imports/<name>.csv
    2. python import_plan.py stage plans/plan_builder.xlsx imports/<name>.csv
         validates the CSV, turns pace keys into targets, matches each workout against
         the Library, and writes Import and Import Schedule tabs for review
    3. python import_plan.py apply plans/plan_builder.xlsx --plan <name> --week-end sun
         after you've filled any blank 'action' cells on the Import tab: adds new
         workouts to Library and saves the plan under <name> (other plans untouched)
       or python import_plan.py apply plans/plan_builder.xlsx --library-only
         to add the new workouts to Library and change nothing else

Close the workbook in Excel before each step.
"""

import argparse
import csv
import difflib
import io
import re
import sys
from pathlib import Path

from openpyxl.styles import Font, PatternFill

from planlib import (
    DAYS,
    PlanError,
    blank,
    day_index,
    last_used_row,
    load_workbook_tables,
    read_table,
    save_workbook,
)

HERE = Path(__file__).parent
TEMPLATE = HERE / "prompts" / "convert_plan.md"
CSV_COLUMNS = ["week", "day", "workout_id", "category", "name", "workout_text", "note"]
CATEGORIES = {"easy", "long", "steady", "tempo", "threshold", "interval", "fartlek",
              "hills", "tuneup", "race"}
OFFSET_KEYS = {"E", "LONG", "STEADY", "M", "HM", "T", "10K", "5K", "I", "R"}
LINE_BREAK = " / "

KEY_RE = re.compile(r"\{([A-Z0-9]+)(?:([+-])(\d+)s)?\}")
LENGTH_RE = re.compile(r"\d+(\.\d+)?(mtr|mi|km|h|m|s)\b")
PCT_RE = re.compile(r"(\d+)% Pace")

IMPORT_SHEET = "Import"
IMPORT_SCHEDULE_SHEET = "Import Schedule"
IMPORT_HEADERS = ["incoming_id", "category", "name", "workout_text", "match",
                  "existing_id", "reason", "differences", "action", "warnings"]


class ImportError_(Exception):
    """Raised when an import can't go ahead. Message lists every problem found."""


# ---------- paces ----------

def load_paces(sheets):
    if "paces" not in sheets:
        raise PlanError("No Paces tab: run setup_workbook.py first")
    return {str(r["key"]).strip().upper(): str(r["target"]).strip() for r in sheets["paces"]}


def parse_pace(text):
    """'7:30/mi' -> (450, 'mi')."""
    m = re.fullmatch(r"\s*(\d+):(\d{2})\s*/\s*(mi|km)\s*", str(text or ""))
    if not m:
        raise PlanError(f"threshold_pace {text!r} should look like 7:30/mi or 4:40/km")
    return int(m[1]) * 60 + int(m[2]), m[3]


def resolve_targets(text, paces, threshold):
    """Replace {KEY} and {KEY+20s} with targets. Returns (text, errors).

    Offsets are seconds of pace per threshold unit. Intervals.icu percentages are % of
    speed, so: anchor pace = T / pct; new pace = anchor pace + offset; new pct = T / new pace.
    """
    errors = []

    def sub(m):
        key, sign, secs = m[1], m[2], m[3]
        if key not in paces:
            errors.append(f"unknown pace key {{{key}}}")
            return m[0]
        target = paces[key]
        if not sign:
            return target
        if key not in OFFSET_KEYS:
            errors.append(f"{{{m[0][1:-1]}}}: offsets don't work on {key}")
            return m[0]
        if threshold is None:
            errors.append(f"{m[0]} needs threshold_pace in Config")
            return m[0]
        pct = int(PCT_RE.fullmatch(target)[1]) / 100
        t_secs = threshold[0]
        new_pace = t_secs / pct + (int(secs) if sign == "+" else -int(secs))
        return f"{round(100 * t_secs / new_pace)}% Pace"

    return KEY_RE.sub(sub, text), errors


# ---------- validation ----------

def steps_and_sections(text):
    """Yield (section_header_or_None, step_line) for each step in ' / ' text."""
    section = None
    for part in (p.strip() for p in text.split(LINE_BREAK)):
        if not part:
            continue
        if part.startswith("-"):
            yield section, part
        else:
            section = part


def check_workout(text, category, raw_text=""):
    """Return (errors, warnings) for one resolved workout_text."""
    errors, warnings = [], []
    steps = list(steps_and_sections(text))
    if not steps:
        return [f"no steps (every step starts with '- ')"], warnings
    for section, step in steps:
        body = step[1:].strip()
        if not step.startswith("- "):
            errors.append(f"step needs a space after '-': {step!r}")
        press = body.lower().startswith("press lap")
        if not press and not LENGTH_RE.search(body):
            errors.append(f"step has no duration or distance: {step!r}")
        sec = (section or "").lower()
        if (sec.startswith("warm") or sec.startswith("cool")) and not press:
            errors.append(f"{section} step must use Press lap: {step!r}")
        if "pace" not in body.lower() and category != "hills" and not body.lower().startswith("uphill"):
            warnings.append(f"step has no pace target: {step!r}")
    if PCT_RE.search(raw_text) or re.search(r"\bZ\d Pace", raw_text):
        warnings.append("uses a raw target instead of a {KEY} from the Paces table")
    if len(steps) == 1 and not steps[0][1][1:].strip().lower().startswith("press lap"):
        errors.append("single-step workouts must use Press lap")
    return errors, sorted(set(warnings))


def normalize(text):
    return re.sub(r"\s+", " ", LINE_BREAK.join(
        p.strip() for p in str(text).split(LINE_BREAK) if p.strip())).strip().lower()


# ---------- matching ----------

def shape(text):
    """Workout structure with lengths, targets and counts taken out:
    '6x / 800mtr 110% / 2m' and '8x / 400mtr 107% / 90s' share a shape."""
    s = normalize(text)
    s = re.sub(r"\d+(\.\d+)?(mtr|mi|km|h|m|s)\b", "LEN", s)
    s = re.sub(r"\d+% pace", "PCT", s)
    return re.sub(r"\d+", "#", s)


def core(text):
    """The workout without its Warmup and Cooldown sections."""
    keep, section = [], None
    for part in (p.strip() for p in str(text).split(LINE_BREAK) if p.strip()):
        if not part.startswith("-"):
            section = part.lower()
        if section and (section.startswith("warm") or section.startswith("cool")):
            continue
        keep.append(part)
    return normalize(LINE_BREAK.join(keep))


def step_count(text):
    return sum(1 for p in str(text).split(LINE_BREAK) if p.strip().startswith("-"))


def differences(old, new):
    a, b = normalize(old).split(), normalize(new).split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
        if op != "equal":
            out.append(f"{' '.join(a[i1:i2]) or '(none)'} -> {' '.join(b[j1:j2]) or '(none)'}")
    return "; ".join(out)


def match_workout(text, category, library):
    """Return (match, existing_id, reason).

    exact:   identical to a Library workout (after tidying whitespace and case)
    similar: a multi-step workout in the same category whose main set is identical apart
             from warmup/cooldown, or has the same structure with different numbers
             (6x800 @ 110% vs 8x800 @ 107%): a candidate to reuse or adapt
    new:     anything else. Single continuous runs of different lengths are distinct
             workouts, never 'similar'.
    """
    norm = normalize(text)
    for wid, row in library.items():
        if not blank(row.get("workout_text")) and normalize(row["workout_text"]) == norm:
            return "exact", wid, "identical"
    if step_count(text) > 1:
        candidates = [(wid, row["workout_text"]) for wid, row in library.items()
                      if not blank(row.get("workout_text"))
                      and str(row.get("category") or "").strip().lower() == category
                      and step_count(row["workout_text"]) > 1]
        for wid, lib_text in candidates:
            if core(lib_text) == core(text):
                return "similar", wid, "same main set, different warmup/cooldown"
        same_shape = [(difflib.SequenceMatcher(None, normalize(t).split(), norm.split()).ratio(), wid)
                      for wid, t in candidates if shape(core(t)) == shape(core(text))]
        if same_shape:
            return "similar", max(same_shape)[1], "same structure, different numbers"
    return "new", None, ""


# ---------- stage ----------

def parse_csv(text):
    """CSV text -> rows. Takes a whole chat reply: when there's a ``` code block, only its
    contents are read, so the notes Claude writes before it don't matter."""
    lines = text.lstrip("\ufeff").splitlines()
    fences = [i for i, ln in enumerate(lines) if ln.strip().startswith("```")]
    if fences:
        lines = lines[fences[0] + 1:fences[1] if len(fences) > 1 else len(lines)]
    reader = csv.DictReader(io.StringIO("\n".join(lines)))
    missing = [c for c in CSV_COLUMNS if c not in (reader.fieldnames or [])]
    if missing:
        raise ImportError_(f"CSV is missing columns: {', '.join(missing)}")
    return list(reader)


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return parse_csv(f.read())


def stage_plan(rows, library, paces, threshold):
    """Validate CSV rows. Returns (workouts, schedule). Raises ImportError_ with all problems."""
    errors, workouts, schedule, seen_days = [], {}, [], set()
    by_norm = {}

    for n, row in enumerate(rows, start=2):
        where = f"CSV line {n}"
        try:
            week = int(str(row["week"]).strip())
            day = DAYS[day_index(row["day"])].title()
        except (ValueError, PlanError) as e:
            errors.append(f"{where}: bad week/day ({e})")
            continue
        if (week, day) in seen_days:
            errors.append(f"{where}: week {week} {day} appears twice")
            continue
        seen_days.add((week, day))

        wid = (row.get("workout_id") or "").strip()
        note = (row.get("note") or "").strip() or None
        if not wid:
            schedule.append({"week": week, "day": day, "incoming_id": None, "note": note})
            continue

        category = (row.get("category") or "").strip().lower()
        if category not in CATEGORIES:
            errors.append(f"{where}: category {category!r} not one of {', '.join(sorted(CATEGORIES))}")
            continue
        raw = (row.get("workout_text") or "").strip()
        if category == "race":
            text, problems, warnings = "", [], []
        else:
            text, problems = resolve_targets(raw, paces, threshold)
            if not problems:
                more, warnings = check_workout(text, category, raw)
                problems += more
            else:
                warnings = []
        if problems:
            errors.extend(f"{where} ({wid}): {p}" for p in problems)
            continue

        norm = normalize(text)
        if wid in workouts and workouts[wid]["norm"] != norm:
            errors.append(f"{where}: {wid} used for two different workouts")
            continue
        if wid not in workouts and norm in by_norm and category != "race":
            wid = by_norm[norm]  # same workout under another id: fold into the first
        if wid not in workouts:
            workouts[wid] = {"incoming_id": wid, "category": category,
                             "name": (row.get("name") or wid).strip(), "workout_text": text,
                             "norm": norm, "warnings": warnings}
            by_norm[norm] = wid
        schedule.append({"week": week, "day": day, "incoming_id": wid, "note": note})

    if errors:
        raise ImportError_("\n".join(errors))

    for w in workouts.values():
        if w["category"] == "race":
            exact = next((i for i, r in library.items()
                          if str(r.get("category") or "").strip().lower() == "race"), None)
            w.update(match="exact" if exact else "new", existing_id=exact, reason="")
        else:
            w["match"], w["existing_id"], w["reason"] = match_workout(
                w["workout_text"], w["category"], library)
        if w["match"] == "exact":
            w["warnings"] = [x for x in w["warnings"] if not x.startswith("uses a raw target")]
        w["differences"] = (differences(library[w["existing_id"]]["workout_text"], w["workout_text"])
                            if w["match"] == "similar" else "")
        w["action"] = {"exact": w["existing_id"], "new": "add", "similar": ""}[w["match"]]
    return list(workouts.values()), schedule


def write_sheet(wb, title, headers, rows, fills=None, widths=None):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", bold=c.row == 1)
    for i, fill in enumerate(fills or [], start=2):
        if fill:
            for c in ws[i]:
                c.fill = fill
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


# ---------- apply ----------

def apply_import(library, workouts, schedule):
    """Resolve actions. Returns (new_library_rows, schedule_rows). Raises ImportError_."""
    errors, final_id, new_rows = [], {}, []
    taken = set(library)
    for w in workouts:
        action = str(w.get("action") or "").strip()
        wid = str(w["incoming_id"]).strip()
        if not action:
            errors.append(f"{wid}: action is blank (enter 'add' or an existing workout_id)")
        elif action.lower() == "add":
            new_id, k = wid, 2
            while new_id in taken:
                new_id, k = f"{wid}-{k}", k + 1
            taken.add(new_id)
            final_id[wid] = new_id
            new_rows.append({"workout_id": new_id, "category": w["category"], "name": w["name"],
                             "workout_text": w["workout_text"] or None})
        elif action in library:
            final_id[wid] = action
        else:
            errors.append(f"{wid}: action {action!r} is not 'add' or a workout_id in Library")
    if errors:
        raise ImportError_("\n".join(errors))
    sched = [{"week": s["week"], "day": s["day"],
              "workout_id": final_id.get(s["incoming_id"]) if s.get("incoming_id") else None,
              "note": s.get("note")} for s in schedule]
    return new_rows, sched


# ---------- commands ----------

def library_listing(library):
    lines = [f"- {wid} ({r.get('category')}): {r.get('workout_text') or ''}" for wid, r in library.items()]
    return "\n".join(lines) or "(empty)"


def build_prompt(config, library, sheets):
    """The conversion prompt with this workbook's paces, offset unit and Library filled in."""
    paces = load_paces(sheets)
    unit = "mile"
    if not blank(config.get("threshold_pace")):
        unit = {"mi": "mile", "km": "km"}[parse_pace(config["threshold_pace"])[1]]
    descriptions = {str(r["key"]).strip().upper(): r.get("description") for r in sheets["paces"]}
    pace_lines = "\n".join(f"- {{{k}}}: {descriptions.get(k) or ''}" for k in paces)
    return (TEMPLATE.read_text(encoding="utf-8")
            .replace("{{PACES}}", pace_lines)
            .replace("{{OFFSET_UNIT}}", unit)
            .replace("{{LIBRARY}}", library_listing(library)))


def cmd_prompt(args):
    _, config, library, sheets = load_workbook_tables(args.workbook)
    text = build_prompt(config, library, sheets)
    out = HERE / "imports" / "prompt.md"
    out.parent.mkdir(exist_ok=True)
    out.write_text(text, encoding="utf-8")
    print(f"Wrote {out}. Paste it into Claude along with the plan, then save the CSV it returns.")
    return 0


def cmd_stage(args):
    wb, config, library, sheets = load_workbook_tables(args.workbook)
    paces = load_paces(sheets)
    threshold = None if blank(config.get("threshold_pace")) else parse_pace(config["threshold_pace"])
    workouts, schedule = stage_plan(read_csv(args.csv), library, paces, threshold)

    yellow, green = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="E2EFDA")
    for w in workouts:
        w["warnings"] = "; ".join(w["warnings"])
    write_sheet(wb, IMPORT_SHEET, IMPORT_HEADERS, workouts,
                fills=[yellow if w["match"] == "similar" else green if w["match"] == "exact" else None
                       for w in workouts],
                widths={"A": 16, "B": 10, "C": 20, "D": 70, "E": 8, "F": 16, "G": 10,
                        "H": 50, "I": 16, "J": 40})
    write_sheet(wb, IMPORT_SCHEDULE_SHEET, ["week", "day", "incoming_id", "note"], schedule,
                widths={"A": 6, "B": 6, "C": 16, "D": 40})
    save_workbook(wb, args.workbook)

    counts = {m: sum(1 for w in workouts if w["match"] == m) for m in ("exact", "similar", "new")}
    print(f"Staged {len(schedule)} days, {len(workouts)} distinct workouts: "
          f"{counts['exact']} already in Library, {counts['new']} new, "
          f"{counts['similar']} similar to an existing workout.")
    for w in workouts:
        if w["match"] == "similar":
            print(f"  SIMILAR {w['incoming_id']} ~ {w['existing_id']} ({w['reason']}): {w['differences']}")
    for w in workouts:
        if w["warnings"]:
            print(f"  warning {w['incoming_id']}: {w['warnings']}")
    if counts["similar"]:
        print("Fill the yellow 'action' cells on the Import tab ('add' or an existing workout_id), then run apply.")
    else:
        print("Review the Import tab, then run apply.")
    return 0


def append_row(ws, values):
    r = last_used_row(ws) + 1
    for c, v in enumerate(values, start=1):
        ws.cell(r, c, v).font = Font(name="Arial")


def cmd_apply(args):
    wb, config, library, sheets = load_workbook_tables(args.workbook)
    if "import" not in sheets or "import schedule" not in sheets:
        raise ImportError_("No Import tabs: run stage first")
    if not args.library_only:
        if not args.plan or not args.week_end:
            raise ImportError_("--plan and --week-end are required (or use --library-only)")
        day_index(args.week_end)
        for needed in ("Plans",):
            if needed not in wb.sheetnames:
                raise ImportError_("No Plans tab: run setup_workbook.py first")
    new_rows, sched = apply_import(library, sheets["import"], sheets["import schedule"])
    source = f"imported from {Path(args.csv_name).name}" if args.csv_name else None
    apply_to_workbook(wb, new_rows, sched, None if args.library_only else args.plan,
                      args.week_end, source)
    save_workbook(wb, args.workbook)

    added = ", ".join(r["workout_id"] for r in new_rows) or "none"
    print(f"Added {len(new_rows)} workouts to Library: {added}")
    if args.library_only:
        print("Schedule, Plans, Config and Layout left as they were.")
        return 0
    plan, week_end = args.plan.strip(), DAYS[day_index(args.week_end)]
    print(f"Plan {plan!r} saved ({len(sched)} days, weeks end {week_end.title()}); "
          f"Config plan set to {plan!r}. Other plans on Schedule are untouched.")
    print("Check race_date and plan_tag in Config, then run layout.py.")
    return 0


def apply_to_workbook(wb, new_rows, sched, plan=None, week_end=None, source=None, select=True):
    """Write an applied import into wb (doesn't save). plan=None adds to Library only.

    select makes the plan Config's plan and removes the Layout and Options tabs, which belonged
    to the previously selected plan.
    """
    lib = wb["Library"]
    for r in new_rows:
        append_row(lib, [r["workout_id"], r["category"], r["name"], r["workout_text"]])
    removed = [IMPORT_SHEET, IMPORT_SCHEDULE_SHEET]
    if plan is not None:
        plan = plan.strip()
        save_plan_rows(wb, plan, DAYS[day_index(week_end)], sched, source)
        if select:
            set_config(wb, "plan", plan)
            removed += ["Layout", "Options"]
    for title in removed:
        if title in wb.sheetnames:
            del wb[title]


def save_plan_rows(wb, plan, week_end, sched, source):
    """Replace this plan's Schedule rows (other plans untouched) and add or update its Plans row."""
    ws = wb["Schedule"]
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    if "plan" not in headers:
        raise ImportError_("Schedule has no plan column: run setup_workbook.py first")
    keep = [r for r in read_table(ws) if str(r.get("plan") or "").strip() != plan]
    ws.delete_rows(2, max(ws.max_row - 1, 1))
    for r in keep:
        append_row(ws, [r.get(h) for h in headers])
    col = {h: i for i, h in enumerate(headers)}
    for s in sched:
        row = [None] * len(headers)
        row[col["plan"]], row[col["week"]], row[col["day"]] = plan, s["week"], s["day"]
        row[col["workout_id"]] = s["workout_id"]
        if "note" in col:
            row[col["note"]] = s["note"]
        append_row(ws, row)

    plans = wb["Plans"]
    found = False
    for r in range(2, last_used_row(plans) + 1):
        if str(plans.cell(r, 1).value or "").strip() == plan:
            plans.cell(r, 2).value = week_end
            found = True
    if not found:
        append_row(plans, [plan, week_end, source])


def set_config(wb, key, value):
    cfg = wb["Config"]
    keys = {str(cfg.cell(r, 1).value).strip(): r for r in range(2, last_used_row(cfg) + 1)
            if cfg.cell(r, 1).value}
    if key in keys:
        cfg.cell(keys[key], 2).value = value
    else:
        append_row(cfg, [key, value])


def main(argv=None):
    parser = argparse.ArgumentParser(description="Import a training plan into plan_builder.xlsx.")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prompt", help="write imports/prompt.md for Claude")
    p.add_argument("workbook")
    s = sub.add_parser("stage", help="validate a CSV and write Import tabs for review")
    s.add_argument("workbook")
    s.add_argument("csv")
    a = sub.add_parser("apply", help="merge the reviewed import into Library and Schedule")
    a.add_argument("workbook")
    a.add_argument("--plan", help="name to save the plan under, e.g. higdon-half-advanced")
    a.add_argument("--week-end", help="weekday the plan's weeks end on, e.g. sun")
    a.add_argument("--csv-name", help=argparse.SUPPRESS)
    a.add_argument("--library-only", action="store_true",
                   help="only add new workouts to Library; leave Schedule, Config and Layout alone")
    args = parser.parse_args(argv)

    try:
        return {"prompt": cmd_prompt, "stage": cmd_stage, "apply": cmd_apply}[args.command](args)
    except (ImportError_, PlanError) as e:
        print(f"Import stopped:\n{e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
