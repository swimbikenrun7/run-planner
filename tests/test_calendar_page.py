"""Drives the Calendar page headlessly. Intervals.icu is always a fake."""

from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import push

ROOT = Path(__file__).resolve().parent.parent
MONDAY = date.today() - timedelta(days=date.today().weekday())
RACE_A = MONDAY + timedelta(days=13)  # plan a: two weeks, race on the second Sunday
B_START = MONDAY + timedelta(days=42)


class FakeIntervals:
    def list_events(self, oldest, newest):
        days = [("a", MONDAY + timedelta(days=k)) for k in range(13)]
        days += [("b", B_START + timedelta(days=k)) for k in range(14)]
        return [{"start_date_local": f"{d}T00:00:00", "name": f"{t} run", "category": "WORKOUT",
                 "external_id": f"{t}-{d}"} for t, d in days]


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    wb.active.title = "Config"
    for name, rows in {
        "Config": [["key", "value"], ["plan", "plan-b"], ["plan_tag", "b"], ["race_date", B_START + timedelta(days=14)]],
        "Library": [["workout_id", "category", "name", "workout_text"]],
        "Plans": [["plan", "source_week_end", "notes", "plan_tag", "race_date", "pushed"],
                  ["plan-a", "sun", None, "a", RACE_A, date.today()], ["plan-b", "sun", None, None, None, None]],
    }.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    path = tmp_path / "wb.xlsx"
    wb.save(path)
    monkeypatch.setenv("RUN_PLANNER_WORKBOOK", str(path))
    monkeypatch.setattr(push, "client_from_env", lambda: FakeIntervals())
    return path


def page():
    return AppTest.from_file(str(ROOT / "ui" / "calendar_page.py"), default_timeout=30).run()


def test_plans_races_and_gaps(workbook):
    at = page()
    assert not at.exception
    plans = at.dataframe[0].value
    assert list(plans["plan_tag"]) == ["a", "b"] and list(plans["plan"]) == ["plan-a", "plan-b"]
    assert list(plans["workouts"]) == [13, 14]
    gap = next(m.value for m in at.markdown if "with nothing planned" in m.value)
    first, last = RACE_A + timedelta(days=1), B_START - timedelta(days=1)
    assert f"{first:%a %b %d, %Y} to {last:%a %b %d, %Y}" in gap and f"{(last - first).days + 1} days" in gap
    weeks = at.dataframe[1].value
    assert "🏁 Race day" in list(weeks.iloc[1])  # plan a's race, second week


def test_without_credentials_it_says_so(workbook, monkeypatch):
    monkeypatch.setattr(push, "client_from_env", lambda: None)
    at = page()
    assert "INTERVALS_API_KEY" in at.error[0].value
