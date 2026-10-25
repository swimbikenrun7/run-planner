import openpyxl
import pytest

from import_plan import (
    ImportError_,
    apply_import,
    apply_to_workbook,
    check_workout,
    match_workout,
    parse_csv,
    resolve_targets,
    stage_plan,
)

PACES = {"RECOVERY": "Z1 Pace", "WU": "Z2 Pace", "E": "80% Pace", "LONG": "80% Pace",
         "M": "94% Pace", "T": "100% Pace", "10K": "103% Pace", "5K": "107% Pace",
         "I": "110% Pace", "R": "118% Pace"}
THRESHOLD = (450, "mi")  # 7:30/mi
WU, CD = "Warmup / - Press lap {WU} 2mi", "Cooldown / - Press lap {WU} 2mi"


def row(week, day, wid="", category="", text="", note=""):
    return {"week": str(week), "day": day, "workout_id": wid, "category": category,
            "name": wid, "workout_text": text, "note": note}


# --- pace keys and offsets

def test_keys_become_targets():
    text, errors = resolve_targets("- Press lap {E} 4mi", PACES, THRESHOLD)
    assert text == "- Press lap 80% Pace 4mi" and not errors


def test_offset_is_seconds_of_pace_converted_to_percent_of_speed():
    # T = 7:30/mi; T+20s = 7:50/mi; 450/470 = 95.7% -> 96%
    assert resolve_targets("- 8mi {T+20s}", PACES, THRESHOLD)[0] == "- 8mi 96% Pace"
    # 10K = 103% -> 7:16.9/mi; +5s -> 7:21.9; 450/441.9 = 101.8% -> 102%
    assert resolve_targets("- 1km {10K+5s}", PACES, THRESHOLD)[0] == "- 1km 102% Pace"


def test_offset_needs_threshold_pace_and_a_percent_key():
    assert resolve_targets("- 8mi {T+20s}", PACES, None)[1]
    assert resolve_targets("- 2m {RECOVERY+10s}", PACES, THRESHOLD)[1]


def test_unknown_key_is_an_error():
    assert resolve_targets("- 1mi {TEMPO}", PACES, THRESHOLD)[1]


# --- the Press lap convention

def test_single_step_must_press_lap():
    assert check_workout("- 30m 80% Pace", "easy")[0]
    assert not check_workout("- Press lap 80% Pace 30m", "easy")[0]


def test_warmup_and_cooldown_must_press_lap():
    bad = "Warmup / - 10m Z2 Pace / Main Set 5x / - 1mi 100% Pace / - 60s Z1 Pace / Cooldown / - Press lap Z2 Pace 10m"
    errors, _ = check_workout(bad, "tempo")
    assert any("Warmup" in e for e in errors)


def test_untimed_recovery_ok_with_press_lap_but_work_needs_length():
    ok = "Hills 5x / - Uphill hard 1m / - Press lap Jog down Z1 Pace"
    assert not check_workout(ok, "hills")[0]
    assert check_workout("Main Set 5x / - 100% Pace / - 60s Z1 Pace", "tempo")[0]


# --- matching against the Library

LIBRARY = {
    "INT-6x800": {"category": "interval", "workout_text":
                  "Warmup / - Press lap Z2 Pace 10m / Main Set 6x / - 800mtr 110% Pace / - 2m Z1 Pace / Cooldown / - Press lap Z2 Pace 10m"},
    "EASY-30": {"category": "easy", "workout_text": "- Press lap 80% Pace 30m"},
}


def test_exact_match_ignores_spacing_and_case():
    assert match_workout("-  press lap 80% pace 30m", "easy", LIBRARY)[:2] == ("exact", "EASY-30")


def test_same_main_set_different_warmup_is_similar():
    text = "Warmup / - Press lap Z2 Pace 2mi / Main Set 6x / - 800mtr 110% Pace / - 2m Z1 Pace / Cooldown / - Press lap Z2 Pace 2mi"
    assert match_workout(text, "interval", LIBRARY)[:2] == ("similar", "INT-6x800")


def test_same_structure_different_numbers_is_similar():
    text = "Warmup / - Press lap Z2 Pace 10m / Main Set 8x / - 400mtr 107% Pace / - 90s Z1 Pace / Cooldown / - Press lap Z2 Pace 10m"
    assert match_workout(text, "interval", LIBRARY)[0] == "similar"


def test_continuous_runs_of_other_lengths_are_new_not_similar():
    assert match_workout("- Press lap 80% Pace 4mi", "easy", LIBRARY)[0] == "new"


# --- staging and applying

def test_stage_collects_every_problem_at_once():
    rows = [row(1, "Mon", "A", "easy", "- 30m {E}"), row(1, "Tue", "B", "jogging", "- Press lap {E} 3mi"),
            row(1, "Funday", "C", "easy", "- Press lap {E} 3mi")]
    with pytest.raises(ImportError_) as e:
        stage_plan(rows, LIBRARY, PACES, THRESHOLD)
    assert str(e.value).count("CSV line") == 3


def test_stage_folds_duplicate_workouts_and_keeps_rest_days():
    rows = [row(1, "Mon", "EASY-4MI", "easy", "- Press lap {E} 4mi"),
            row(1, "Tue", note="Rest or cross-train"),
            row(1, "Wed", "EASY-FOUR", "easy", "- Press lap {E} 4mi")]
    workouts, schedule = stage_plan(rows, LIBRARY, PACES, THRESHOLD)
    assert [w["incoming_id"] for w in workouts] == ["EASY-4MI"]
    assert [s["incoming_id"] for s in schedule] == ["EASY-4MI", None, "EASY-4MI"]


def test_stage_rejects_one_id_for_two_workouts():
    rows = [row(1, "Mon", "EASY", "easy", "- Press lap {E} 4mi"),
            row(1, "Tue", "EASY", "easy", "- Press lap {E} 5mi")]
    with pytest.raises(ImportError_, match="two different workouts"):
        stage_plan(rows, LIBRARY, PACES, THRESHOLD)


def test_apply_reuses_adds_and_never_overwrites():
    workouts = [
        {"incoming_id": "EASY-30", "category": "easy", "name": "x", "workout_text": "- Press lap 80% Pace 31m", "action": "add"},
        {"incoming_id": "INT-NEW", "category": "interval", "name": "y", "workout_text": "...", "action": "INT-6x800"},
    ]
    schedule = [{"week": 1, "day": "Mon", "incoming_id": "EASY-30", "note": None},
                {"week": 1, "day": "Tue", "incoming_id": "INT-NEW", "note": None},
                {"week": 1, "day": "Wed", "incoming_id": None, "note": "Rest"}]
    new_rows, sched = apply_import(LIBRARY, workouts, schedule)
    assert [r["workout_id"] for r in new_rows] == ["EASY-30-2"]  # renamed, not overwritten
    assert [s["workout_id"] for s in sched] == ["EASY-30-2", "INT-6x800", None]


def test_apply_requires_every_action():
    with pytest.raises(ImportError_, match="blank"):
        apply_import(LIBRARY, [{"incoming_id": "X", "category": "easy", "name": "x",
                                "workout_text": "t", "action": None}], [])


# --- pasted CSV and saving to the workbook

def test_pasted_chat_reply_reads_only_the_csv_block():
    text = ("Interpreted 'hard' as I pace.\n\n```csv\nweek,day,workout_id,category,name,workout_text,note\n"
            "1,Mon,EASY-3MI,easy,Easy 3mi,- Press lap {E} 3mi,\n```\nLet me know if...\n")
    assert [r["workout_id"] for r in parse_csv(text)] == ["EASY-3MI"]
    with pytest.raises(ImportError_, match="missing columns"):
        parse_csv("week,day\n1,Mon\n")


def workbook(config_plan="old"):
    wb = openpyxl.Workbook()
    wb.active.title = "Config"
    for name, rows in {"Config": [["key", "value"], ["plan", config_plan]],
                       "Library": [["workout_id", "category", "name", "workout_text"]],
                       "Paces": [["key", "description", "target"], ["E", "", "80% Pace"]],
                       "Plans": [["plan", "source_week_end", "notes", "rest_fill", "pre_race_fill"],
                                 ["old", "sun", None, None, None]],
                       "Schedule": [["plan", "week", "day", "workout_id", "note"], ["old", 1, "Mon", "X", None]],
                       "Layout": [["date"]]}.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    return wb


NEW = [{"workout_id": "E3", "category": "easy", "name": "Easy", "workout_text": "- Press lap 80% Pace 3mi"}]
SCHED = [{"week": 1, "day": "Mon", "workout_id": "E3", "note": None}]


def test_saving_a_plan_without_selecting_it_keeps_config_and_layout():
    wb = workbook()
    apply_to_workbook(wb, NEW, SCHED, "new", "sun", select=False)
    assert wb["Config"]["B2"].value == "old" and "Layout" in wb.sheetnames
    assert [r[0] for r in wb["Schedule"].iter_rows(min_row=2, values_only=True)] == ["old", "new"]
    assert wb["Library"]["A2"].value == "E3"


def test_saving_and_selecting_a_plan_replaces_the_old_layout():
    wb = workbook()
    apply_to_workbook(wb, NEW, SCHED, "new", "sun")
    assert wb["Config"]["B2"].value == "new" and "Layout" not in wb.sheetnames


def plans_row(wb, plan):
    return next(r for r in wb["Plans"].iter_rows(values_only=True) if r[0] == plan)


def test_new_plan_gets_rest_fill_in_its_own_units():
    wb = workbook()
    apply_to_workbook(wb, NEW, SCHED, "miles", "sun", select=False)  # E3 is written in miles
    assert plans_row(wb, "miles")[3:5] == ("EASY-3MI", "EASY-2MI")
    by_time = [{**NEW[0], "workout_id": "E30", "workout_text": "- Press lap 80% Pace 30m"}]
    apply_to_workbook(wb, by_time, [{**SCHED[0], "workout_id": "E30"}], "minutes", "sun", select=False)
    assert plans_row(wb, "minutes")[3:5] == ("EASY-30", "EASY-20")
    library = {r[0]: r for r in wb["Library"].iter_rows(min_row=2, values_only=True)}
    assert library["EASY-20"][3] == "- Press lap 80% Pace 20m"  # added, from the Paces E target


def test_reimporting_a_plan_keeps_the_users_rest_fill_choice():
    wb = workbook()
    wb["Plans"]["D2"].value = wb["Plans"]["E2"].value = None  # user cleared: keep rest days
    apply_to_workbook(wb, NEW, SCHED, "old", "sun", select=False)
    assert plans_row(wb, "old")[3:5] == (None, None)
