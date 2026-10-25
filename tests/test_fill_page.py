"""Drives the Fill gaps page headlessly. Intervals.icu is always a fake."""

from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import push

ROOT = Path(__file__).resolve().parent.parent
TODAY = date.today()
RACE_A = TODAY + timedelta(days=10)  # plan a ends with a 10K
B_START = RACE_A + timedelta(days=50)  # plan b starts after a seven-week gap
EASY = "- Press lap 80% Pace 30m"


class FakeIntervals:
    def __init__(self, extra=()):
        self.upserted, self.deleted, self.extra = None, None, list(extra)

    def list_events(self, oldest, newest):
        days = [("a", TODAY + timedelta(days=k)) for k in range(10)]
        days += [("b", B_START + timedelta(days=k)) for k in range(21)]
        return [{"start_date_local": f"{d}T00:00:00", "name": "Easy 30", "category": "WORKOUT",
                 "external_id": f"{t}-{d}", "icu_training_load": 32} for t, d in days] + self.extra

    def wellness(self, oldest, newest):
        return [{"id": (TODAY - timedelta(days=k)).isoformat(), "ctl": 30.0, "atl": 35.0} for k in range(4)]

    def run_settings(self):
        return {"threshold_speed": 3.8779278, "pace_zones": [77.5, 87.7, 94.3, 100.0, 103.4, 111.5, 999.0]}

    def upsert(self, events):
        self.upserted = events

    def delete(self, external_ids):
        self.deleted = external_ids


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    wb.active.title = "Config"
    days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    roles = {"Mon": "quality", "Wed": "quality", "Fri": "quality", "Sat": "long"}
    for name, rows in {
        "Config": [["key", "value"], ["plan", "pb"], ["plan_tag", "b"]],
        "Paces": [["key", "description", "target"], ["E", "", "80% Pace"], ["LONG", "", "80% Pace"],
                  ["10K", "", "103% Pace"], ["HM", "", "98% Pace"]],
        "Library": [["workout_id", "category", "name", "workout_text"],
                    ["EASY-30", "easy", "Easy 30", EASY],
                    ["TEMPO", "tempo", "Tempo", "Warmup / - Press lap Z2 Pace 10m / Main Set / - 20m 100% Pace / "
                                                "Cooldown / - Press lap Z2 Pace 10m"]],
        "Week": [["day", "role", "workout_id"]] + [[d, roles.get(d, "easy"), None] for d in days],
        "Plans": [["plan", "source_week_end", "notes", "rest_fill", "pre_race_fill", "plan_tag", "race_date",
                   "pushed", "race_distance"],
                  ["pa", "sun", None, "EASY-30", None, "a", datetime.combine(RACE_A, datetime.min.time()), None, "10k"],
                  ["pb", "sun", None, "EASY-30", None, "b", None, None, "half"]],
    }.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    path = tmp_path / "wb.xlsx"
    wb.save(path)
    monkeypatch.setenv("RUN_PLANNER_WORKBOOK", str(path))
    return path


def page(monkeypatch, fake):
    monkeypatch.setattr(push, "client_from_env", lambda: fake)
    return AppTest.from_file(str(ROOT / "ui" / "fill_page.py"), default_timeout=60).run()


def test_gap_preview_follows_the_rules(workbook, monkeypatch):
    at = page(monkeypatch, FakeIntervals())
    assert not at.exception, at.exception
    gap = at.selectbox[0].value
    assert gap[0] == RACE_A + timedelta(days=1) and gap[1] == B_START - timedelta(days=1)
    assert "after the 10k race (a)" in at.selectbox[0].format_func(gap) and "before b" in at.selectbox[0].format_func(gap)
    grid = at.dataframe[0].value
    first_days = [v for v in grid.iloc[0] if isinstance(v, str) and v]
    assert first_days[0] == "Easy 20"  # the day after the race: the shortest easy run (time units, like plan b)
    assert any("Tempo" == v for v in grid.values.ravel())  # a quality session from the Library


def test_saving_rules_writes_config(workbook, monkeypatch):
    at = page(monkeypatch, FakeIntervals())
    at.number_input(key="r_10k").set_value(5).run()
    at.radio(key="r_between").set_value("match").run()
    next(b for b in at.button if b.label == "Save rules").click().run()
    config = {r[0]: r[1] for r in openpyxl.load_workbook(workbook)["Config"].iter_rows(values_only=True)}
    assert "10k:5" in config["recovery_days"] and config["between_plans"] == "match"


def test_push_records_the_fill_and_adds_new_workouts(workbook, monkeypatch):
    stale = f"fill-{RACE_A + timedelta(days=1)}"
    fake = FakeIntervals(extra=[{"start_date_local": f"{B_START - timedelta(days=2)}T00:00:00", "name": "old",
                                 "category": "WORKOUT", "external_id": f"{stale}-{B_START - timedelta(days=200)}",
                                 "icu_training_load": 30}])
    at = page(monkeypatch, fake)
    assert next(b for b in at.button if b.label == "Push").disabled  # nothing without confirming
    at.checkbox(key="fill_confirm").check().run()
    next(b for b in at.button if b.label == "Push").click().run()
    assert not at.exception and "Pushed" in at.success[0].value
    assert {e["external_id"].rsplit("-", 3)[0] for e in fake.upserted} == {stale}  # reuses the gap's fill tag
    wb = openpyxl.load_workbook(workbook)
    fill_rows = list(wb["Fill"].iter_rows(min_row=2, values_only=True))
    assert len(fill_rows) == (B_START - RACE_A).days - 1 and {r[0] for r in fill_rows} == {stale}
    fitness, fatigue, form, pushed = fill_rows[0][5:9]  # the forecast, kept to compare with what happens
    assert fitness > 0 and fatigue > 0 and form is not None and pushed.date() == TODAY
    library = {r[0]: r for r in wb["Library"].iter_rows(min_row=2, values_only=True)}
    assert library["EASY-30"][3] == EASY  # existing workouts untouched
    assert all(e["description"] for e in fake.upserted)  # every pushed workout has its text


def test_without_credentials_it_says_so(workbook, monkeypatch):
    at = page(monkeypatch, None)
    assert "INTERVALS_API_KEY" in at.error[0].value


def test_with_no_plan_after_it_a_gap_is_filled_90_days_ahead(workbook, monkeypatch):
    at = page(monkeypatch, FakeIntervals())
    open_gap = at.selectbox[0].options[-1]
    assert "no plan after it yet (fills to" in open_gap
    at.selectbox[0].select_index(len(at.selectbox[0].options) - 1).run()
    assert not at.exception, at.exception
    at.checkbox(key="fill_confirm").check().run()
    fake = FakeIntervals()
    monkeypatch.setattr(push, "client_from_env", lambda: fake)
    next(b for b in at.button if b.label == "Push").click().run()
    last = max(date.fromisoformat(e["start_date_local"][:10]) for e in fake.upserted)
    assert last == TODAY + timedelta(days=90)


class FarIntervals(FakeIntervals):
    """Plan b runs past the 90 days, with nothing after it."""

    def list_events(self, oldest, newest):
        return [e for e in super().list_events(oldest, newest)] + [
            {"start_date_local": f"{TODAY + timedelta(days=k)}T00:00:00", "name": "Easy 30", "category": "WORKOUT",
             "external_id": f"b-{TODAY + timedelta(days=k)}", "icu_training_load": 32} for k in range(80, 101)]


def test_a_gap_beyond_90_days_with_no_plan_after_waits(workbook, monkeypatch):
    at = page(monkeypatch, FarIntervals())
    at.selectbox[0].select_index(len(at.selectbox[0].options) - 1).run()
    assert not at.exception, at.exception
    assert "filled from" in at.info[0].value and not [b for b in at.button if b.label == "Push"]


def test_filler_split_by_a_plan_is_flagged_and_replaced(workbook, monkeypatch):
    old = f"fill-{RACE_A - timedelta(days=20)}"  # pushed before plan a was put on the calendar
    fake = FakeIntervals(extra=[{"start_date_local": f"{RACE_A + timedelta(days=d)}T00:00:00", "name": "old",
                                 "category": "WORKOUT", "external_id": f"{old}-{RACE_A + timedelta(days=d)}",
                                 "icu_training_load": 30} for d in (3, 30)])
    at = page(monkeypatch, fake)
    assert "out of date" in at.selectbox[0].format_func(at.selectbox[0].value)
    assert any("planned before a plan was pushed" in w.value for w in at.warning)
    at.checkbox(key="fill_confirm").check().run()
    next(b for b in at.button if b.label == "Push").click().run()
    assert {e["external_id"][:15] for e in fake.upserted} == {f"fill-{RACE_A + timedelta(days=1)}"}
    assert sorted(fake.deleted) == [f"{old}-{RACE_A + timedelta(days=d)}" for d in (3, 30)]  # under the old tag
