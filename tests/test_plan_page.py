"""Drives the Plan page headlessly with Streamlit's AppTest. Intervals.icu is always a fake."""

from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import push

ROOT = Path(__file__).resolve().parent.parent
_SOON = date.today() + timedelta(days=60)
RACE = _SOON + timedelta(days=6 - _SOON.weekday())  # a Sunday about two months out
WEEK = ["EASY-3MI", "INT", "EASY-3MI", "INT", None, "INT", "LONG"]  # three hard days + long
RACE_WEEK = ["EASY-3MI", "INT", "EASY-2MI", None, None, None, "RACE"]


class FakeIntervals:
    def __init__(self):
        self.upserted = self.deleted = None

    def list_events(self, oldest, newest):
        # one event of this plan after the race (no workout there now: stale), one of another plan
        return [{"external_id": f"t-{RACE + timedelta(days=3)}"}, {"external_id": "other-2030-01-01"}]

    def upsert(self, events):
        self.upserted = events

    def delete(self, external_ids):
        self.deleted = external_ids


def make_workbook(path, week_roles, allow_pre_long="yes"):
    wb = openpyxl.Workbook()
    wb.active.title = "Config"
    days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    tables = {
        "Config": [["key", "value"], ["plan", "p"], ["race_date", RACE], ["plan_tag", "t"],
                   ["allow_pre_long", allow_pre_long]],
        "Library": [["workout_id", "category", "name", "workout_text"],
                    ["EASY-3MI", "easy", "Easy 3mi", "- Press lap 80% Pace 3mi"],
                    ["EASY-2MI", "easy", "Easy 2mi", "- Press lap 80% Pace 2mi"],
                    ["INT", "interval", "Intervals", "- 5m 110% Pace"],
                    ["LONG", "long", "Long", "- Press lap 80% Pace 90m"],
                    ["RACE", "race", "Race", None]],
        "Week": [["day", "role", "workout_id"]] + [[d, week_roles.get(d, "easy"), None] for d in days],
        "Plans": [["plan", "source_week_end", "notes", "rest_fill", "pre_race_fill", "plan_tag", "race_date", "pushed"],
                  ["p", "sun", None, "EASY-3MI", "EASY-2MI", None, None, None]],
        "Schedule": [["plan", "week", "day", "workout_id", "note"]]
                    + [["p", w, d, wid, None] for w, plan_week in ((1, WEEK), (2, RACE_WEEK))
                       for d, wid in zip(days, plan_week)],
    }
    for name, rows in tables.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    path = tmp_path / "wb.xlsx"
    make_workbook(path, {"Tue": "quality", "Thu": "quality", "Sat": "quality", "Sun": "long"})
    monkeypatch.setenv("RUN_PLANNER_WORKBOOK", str(path))
    monkeypatch.setattr(push, "client_from_env", lambda: None)  # never the real calendar
    return path


def page():
    return AppTest.from_file(str(ROOT / "ui" / "plan_page.py"), default_timeout=30).run()


def click(at, label):
    next(b for b in at.button if b.label == label).click().run()
    assert not at.exception, at.exception
    return at


def labelled(widgets, start):
    return next(w for w in widgets if w.label.startswith(start))


def sheet(path, name):
    return [r for r in openpyxl.load_workbook(path)[name].iter_rows(min_row=2, values_only=True) if any(r)]


def test_lay_out_writes_the_layout_and_shows_the_calendar(workbook):
    at = click(page(), "Lay out")
    assert "Layout written" in at.success[0].value
    rows = sheet(workbook, "Layout")
    assert rows[-1][3] == "RACE" and all(r[3] for r in rows)  # rest days filled
    assert at.dataframe  # the calendar


def test_options_appear_as_cards_and_choosing_one_updates_the_week(workbook, monkeypatch):
    make_workbook(workbook, {"Tue": "quality", "Sun": "long"}, "no")  # hard sessions won't fit
    at = click(page(), "Lay out")
    assert "don't fit your week" in at.warning[0].value
    assert "Layout" not in openpyxl.load_workbook(workbook).sheetnames  # advises, writes nothing
    next(b for b in at.button if b.key == "pick1").click().run()
    assert not at.exception and "Week tab was updated" in at.success[0].value
    roles = [r[1] for r in sheet(workbook, "Week")]
    assert roles.count("quality") > 1 and sheet(workbook, "Layout")


def test_save_settings_writes_config_and_plans(workbook):
    at = page()
    at.text_input[0].input("t2").run()
    labelled(at.selectbox, "Rest days become").select("").run()  # keep rest days
    at = click(at, "Save settings")
    config = dict(r[:2] for r in sheet(workbook, "Config"))
    assert config["plan_tag"] == "t2"
    assert sheet(workbook, "Plans")[0][3] is None


def test_check_and_push_use_intervals_only_after_confirming(workbook, monkeypatch):
    fake = FakeIntervals()
    monkeypatch.setattr(push, "client_from_env", lambda: fake)
    at = click(click(page(), "Lay out"), "Check against Intervals")
    assert fake.upserted is None  # the check sends nothing
    assert next(b for b in at.button if b.label == "Push").disabled
    assert any("would be **deleted**" in w.value for w in at.warning)
    at.checkbox(key="confirm").check().run()
    at = click(at, "Push")
    assert "Pushed" in at.success[0].value and fake.upserted and fake.deleted == [f"t-{RACE + timedelta(days=3)}"]
    assert "Couldn't note" not in at.success[0].value
    tag, race, pushed = sheet(workbook, "Plans")[0][5:8]
    assert (tag, race.date(), pushed.date()) == ("t", RACE, date.today())  # recorded for the master calendar


def test_without_credentials_the_check_works_but_push_is_off(workbook):
    at = click(click(page(), "Lay out"), "Check against Intervals")
    assert any("No .env credentials" in w.value for w in at.warning)
    assert not [b for b in at.button if b.label == "Push"]


def test_week_dropdowns_and_rules_are_saved_and_used_by_layout(workbook):
    at = page()
    at.selectbox(key="day3").select("EASY-2MI").run()  # Thursday: club run
    at.selectbox(key="day5").select("long").run()       # Saturday long...
    at.selectbox(key="day6").select("").run()           # ...so Sunday is open
    at.toggle[0].set_value(False).run()
    at = click(at, "Save settings")
    week = {r[0]: (r[1], r[2]) for r in sheet(workbook, "Week")}
    assert week["Thu"] == ("fixed", "EASY-2MI") and week["Sat"] == ("long", None) and week["Sun"] == ("easy", None)
    assert dict(r[:2] for r in sheet(workbook, "Config"))["allow_pre_long"] == "no"
    at = click(at, "Lay out")
    if not at.success:  # the plan doesn't fit this week: keep the week as it is
        at.button(key="pick0").click().run()
    rows = sheet(workbook, "Layout")
    assert [r[3] for r in rows if r[1] == "Thu"] == ["EASY-2MI", "EASY-2MI"]
    assert {r[3] for r in rows if r[1] == "Sat" and r[0]} <= {"LONG", "EASY-2MI"}  # long day (shakeout before race)


def test_unsaved_changes_explain_why_the_check_is_locked(workbook):
    at = page()
    at.text_input[0].input("t2").run()
    assert next(b for b in at.button if b.label == "Check against Intervals").disabled
    assert any("changes that aren't saved" in i.value for i in at.info)


def test_two_long_days_cannot_be_saved(workbook):
    at = page()
    at.selectbox(key="day5").select("long").run()  # Sunday is already long
    assert any("Only one long-run day" in e.value for e in at.error)
    assert next(b for b in at.button if b.label == "Save settings").disabled


def test_after_choosing_an_option_you_can_go_straight_to_the_check(workbook):
    make_workbook(workbook, {"Tue": "quality", "Sun": "long"}, "no")  # hard sessions won't fit
    at = click(page(), "Lay out")
    at.button(key="pick1").click().run()
    assert "Week tab was updated" in at.success[0].value
    saved = {r[0]: r[1] for r in sheet(workbook, "Week")}
    assert [at.selectbox(key=f"day{d}").value or "easy" for d in range(7)] == \
        [saved[d] if saved[d] != "fixed" else None for d in ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]]
    assert next(b for b in at.button if b.label == "Save settings").disabled  # nothing unsaved
    assert not next(b for b in at.button if b.label == "Check against Intervals").disabled


def test_changes_saved_elsewhere_reload_section_one(workbook):
    at = page()
    wb = openpyxl.load_workbook(workbook)  # as if edited and saved in Calc
    wb["Week"]["B3"].value = "long"  # Tue
    wb["Week"]["B8"].value = "easy"  # Sun
    wb.save(workbook)
    at.run()
    assert at.selectbox(key="day1").value == "long" and at.selectbox(key="day6").value == ""
    assert next(b for b in at.button if b.label == "Save settings").disabled  # nothing unsaved


def test_unsaved_changes_are_named(workbook):
    at = page()
    at.selectbox(key="day0").select("quality").run()
    assert any("Not saved yet: Mon" in c.value for c in at.caption)
    assert any("(Mon)" in i.value for i in at.info)
